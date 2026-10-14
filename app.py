import streamlit as st
import pandas as pd
from sqlalchemy import create_engine, text
import os
from io import BytesIO
from pathlib import Path
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv(Path(__file__).resolve().with_name(".env"))


REFUND_SHEET_NAME = "Refund Reconciliation"
REFUND_REPORT_COLUMNS = [
    "Receipt ID",
    "Withdraw Amount",
    "Order Amount",
    "Refund Session ID",
    "Transaction Status",
    "Mismatch Reason",
]


def normalize_receipt_id(value):
    if pd.isna(value):
        return ""
    receipt_id = str(value).strip()
    if receipt_id.endswith(".0") and receipt_id[:-2].isdigit():
        receipt_id = receipt_id[:-2]
    return receipt_id


def build_refund_reconciliation(statement_df, refunded_orders_df):
    statement_df = statement_df.copy()
    refunded_orders_df = refunded_orders_df.copy()

    statement_df["Receipt No."] = statement_df["Receipt No."].map(normalize_receipt_id)
    withdrawal_text = (
        statement_df["Withdrawn"]
        .fillna("")
        .astype(str)
        .str.strip()
    )
    statement_df["Withdrawn"] = pd.to_numeric(
        withdrawal_text.str.replace(",", "", regex=False).str.replace(
            r"[^0-9.\-]", "", regex=True
        ),
        errors="coerce",
    )
    statement_df["_has_withdrawal"] = withdrawal_text != ""
    statement_withdrawals = statement_df[
        statement_df["_has_withdrawal"] & statement_df["Withdrawn"].lt(0)
    ]

    if "refund_session_id" not in refunded_orders_df:
        refunded_orders_df["refund_session_id"] = ""
    for column in ("order_amount", "tip", "adjusment", "payment_status"):
        if column not in refunded_orders_df:
            refunded_orders_df[column] = None
    refunded_orders_df["_receipt_id"] = refunded_orders_df[
        "refund_session_id"
    ].map(normalize_receipt_id)
    refunded_orders_df["_expected_withdrawal"] = (
        pd.to_numeric(refunded_orders_df["order_amount"], errors="coerce")
        + pd.to_numeric(refunded_orders_df["tip"], errors="coerce").fillna(0)
        + pd.to_numeric(refunded_orders_df["adjusment"], errors="coerce").fillna(0)
    ).round(2)

    refunded_by_receipt = {
        receipt_id: group
        for receipt_id, group in refunded_orders_df.groupby("_receipt_id", dropna=False)
        if receipt_id
    }
    statement_receipts = set(
        receipt_id
        for receipt_id in statement_withdrawals["Receipt No."]
        if receipt_id
    )
    reconciled_rows = []

    def make_report_row(statement_record=None, order=None, mismatch_reason=""):
        row = {
            "Receipt ID": "",
            "Withdraw Amount": None,
            "Order Amount": None,
            "Refund Session ID": "",
            "Transaction Status": "",
            "Mismatch Reason": mismatch_reason,
        }
        if statement_record is not None:
            row["Receipt ID"] = statement_record["Receipt No."]
            row["Withdraw Amount"] = statement_record["Withdrawn"]
        if order is not None:
            if not row["Receipt ID"]:
                row["Receipt ID"] = order["_receipt_id"]
            row["Order Amount"] = order.get("order_amount")
            row["Refund Session ID"] = order.get("refund_session_id")
            row["Transaction Status"] = order.get("payment_status")
        return row

    for _, statement_record in statement_withdrawals.iterrows():
        receipt_id = statement_record["Receipt No."]
        matches = refunded_by_receipt.get(receipt_id)

        if not receipt_id:
            reconciled_rows.append(make_report_row(
                statement_record,
                mismatch_reason="Missing receipt ID in statement",
            ))
            continue
        if matches is None:
            reconciled_rows.append(make_report_row(
                statement_record,
                mismatch_reason="Receipt ID not found in refunded orders",
            ))
            continue

        for _, order in matches.iterrows():
            withdrawn = statement_record["Withdrawn"]
            expected_withdrawal = order["_expected_withdrawal"]
            if pd.isna(withdrawn):
                reason = "Withdrawn amount is not numeric"
            elif pd.isna(expected_withdrawal):
                reason = "Order amount is unavailable"
            elif round(abs(float(withdrawn)), 2) != round(abs(float(expected_withdrawal)), 2):
                reason = "Withdraw amount does not equal order amount + tip + adjustment"
            else:
                reason = ""
            reconciled_rows.append(make_report_row(
                statement_record,
                order,
                reason,
            ))

    for _, order in refunded_orders_df.iterrows():
        receipt_id = order["_receipt_id"]
        if not receipt_id or receipt_id not in statement_receipts:
            reason = (
                "Refunded order has no refund receipt ID"
                if not receipt_id
                else "Refund receipt ID not found in statement withdrawals"
            )
            reconciled_rows.append(make_report_row(
                order=order,
                mismatch_reason=reason,
            ))

    report_df = pd.DataFrame(reconciled_rows, columns=REFUND_REPORT_COLUMNS)
    return report_df, statement_withdrawals


def build_excel_report(uploaded_bytes, uploaded_name, report_df, start_date, end_date):
    if uploaded_name.lower().endswith(".xlsx"):
        workbook = load_workbook(BytesIO(uploaded_bytes))
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
        source_excel = pd.ExcelFile(BytesIO(uploaded_bytes))
        for sheet_name in source_excel.sheet_names:
            source_df = pd.read_excel(
                source_excel, sheet_name=sheet_name, header=None
            )
            worksheet = workbook.create_sheet(title=sheet_name[:31])
            for row_number, values in enumerate(
                source_df.itertuples(index=False, name=None), start=1
            ):
                for column_number, value in enumerate(values, start=1):
                    if pd.isna(value):
                        value = None
                    elif hasattr(value, "item"):
                        value = value.item()
                    worksheet.cell(row_number, column_number, value)

    if REFUND_SHEET_NAME in workbook.sheetnames:
        del workbook[REFUND_SHEET_NAME]
    worksheet = workbook.create_sheet(REFUND_SHEET_NAME)
    worksheet["A1"] = "Telebirr Refund Reconciliation"
    worksheet["A2"] = f"Date range: {start_date} to {end_date}"
    for cell_ref in ("A1", "A2"):
        worksheet[cell_ref].font = Font(bold=True)

    header_row = 4
    for column_number, column_name in enumerate(report_df.columns, start=1):
        cell = worksheet.cell(header_row, column_number, column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for row_number, values in enumerate(
        report_df.itertuples(index=False, name=None), start=header_row + 1
    ):
        is_mismatch = bool(values[-1])
        for column_number, value in enumerate(values, start=1):
            if pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            cell = worksheet.cell(row_number, column_number, value)
            if is_mismatch:
                cell.fill = PatternFill("solid", fgColor="FFC7CE")
                cell.font = Font(color="9C0006")

    if len(report_df.columns):
        worksheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(report_df.columns))}"
            f"{header_row + len(report_df)}"
        )
    worksheet.freeze_panes = "A5"
    for column_number, column_name in enumerate(report_df.columns, start=1):
        width = min(max(len(str(column_name)) + 2, 14), 40)
        worksheet.column_dimensions[get_column_letter(column_number)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


# --- Database Connection Setup ---
@st.cache_resource
def init_connection():
    # Use the synchronous MySQL dialect because pd.read_sql() is running synchronously.
    # mysql+aiomysql triggers the async SQLAlchemy stack and requires greenlet.
    db_url = os.getenv("DATABASE_URL")
    if not db_url:
        raise RuntimeError(
            "DATABASE_URL is not set. Add it to payment_check/.env or set it "
            "as an environment variable."
        )
    return create_engine(db_url)

engine = init_connection()

st.title("Telebirr Reconciliation Dashboard")
st.write("Upload a Telebirr statement and select a date range to reconcile it against the transaction and refund records.")

# --- UI Inputs ---
col1, col2 = st.columns(2)
with col1:
    start_date = st.date_input("Start Date")
with col2:
    end_date = st.date_input("End Date")

uploaded_file = st.file_uploader("Upload Telebirr Statement (Excel)", type=["xlsx", "xls"])

if uploaded_file and st.button("Run Reconciliation"):
    try:
        # 1. Read the uploaded Excel file
        # The Telebirr statement has metadata rows above the actual table,
        # and the real column headers begin on row 6 (1-based indexing).
        excel_df = pd.read_excel(uploaded_file, header=5)

        # Normalize column names and support small variations in the header text
        excel_df.columns = [str(col).strip() for col in excel_df.columns]

        if 'Receipt No.' not in excel_df.columns:
            receipt_candidates = [
                c for c in excel_df.columns
                if 'receipt' in str(c).lower() and 'no' in str(c).lower()
            ]
            if receipt_candidates:
                excel_df = excel_df.rename(columns={receipt_candidates[0]: 'Receipt No.'})

        if 'Receipt No.' not in excel_df.columns:
            st.error("The uploaded Excel file must contain a 'Receipt No.' column in row 6.")
            st.stop()

        withdrawn_candidates = [
            column for column in excel_df.columns
            if str(column).strip().lower() == "withdrawn"
        ]
        if not withdrawn_candidates:
            withdrawn_candidates = [
                column for column in excel_df.columns
                if "withdrawn" in str(column).strip().lower()
            ]
        if not withdrawn_candidates:
            st.error("The uploaded Excel file must contain a 'Withdrawn' column in row 6.")
            st.stop()
        if withdrawn_candidates[0] != "Withdrawn":
            excel_df = excel_df.rename(columns={withdrawn_candidates[0]: "Withdrawn"})

        refund_statement_df = excel_df.copy()

        # Clean the Receipt No. column to ensure perfect matching
        excel_df['Receipt No.'] = excel_df['Receipt No.'].fillna('').astype(str).str.strip()
        excel_df = excel_df[excel_df['Receipt No.'] != '']

        # 2. Fetch data from DB based on selected date range and payment method
        # Note: Depending on your DB driver (MySQL vs PostgreSQL), you might need to change %s to ? or %(param)s
        query = """
            SELECT 
                ct.session_id,
                ct.order_id,
                ct.transaction_status,
                o.payment_method,
                ct.created_at
            FROM customer_transactions ct
            LEFT JOIN orders o ON o.id = ct.order_id
            WHERE DATE(ct.created_at) >= %s
              AND DATE(ct.created_at) <= %s
              AND o.payment_method IN ('telebirr-ussd', 'telebirr')
        """
        
        # Fetch dataframe
        db_df = pd.read_sql(query, engine, params=(start_date, end_date))
        
        # Clean session_id for matching
        db_df['session_id'] = db_df['session_id'].astype(str).str.strip()
        
        # Get only the successful transactions for baseline comparison
        db_success_df = db_df[db_df['transaction_status'].str.lower() == 'success']

        # --- RECONCILIATION LOGIC ---

        # Check 1: In Database as success/paid, but NOT in Telebirr Excel Statement
        st.subheader("Check 1: Paid in Database, Missing in Telebirr Statement")
        missing_in_telebirr = db_success_df[~db_success_df['session_id'].isin(excel_df['Receipt No.'])].copy()
        
        # Styling function for red highlights
        def highlight_red(row):
            return ['background-color: #ffcccc; color: #990000; font-weight: bold'] * len(row)

        if not missing_in_telebirr.empty:
            missing_in_telebirr['Mismatch Reason'] = "Doesn't exist in Telebirr statement"
            # Apply styling and display
            st.dataframe(
                missing_in_telebirr.style.apply(highlight_red, axis=1),
                use_container_width=True
            )
        else:
            st.success("All successful telebirr database transactions are present in the Excel statement!")

        # Check 2: In Telebirr Excel Statement, but NOT in Database (or not marked as 'success')
        st.subheader("Check 2: In Telebirr Statement, Missing or Failed in Database")
        missing_in_db = excel_df[~excel_df['Receipt No.'].isin(db_success_df['session_id'])].copy()

        if not missing_in_db.empty:
            missing_in_db['Mismatch Reason'] = "Not in customer transaction or status != success"
            
            # Map DB statuses to the excel view for context (optional but highly useful)
            # Find what the status actually is in the DB, if it exists at all
            db_status_map = db_df.set_index('session_id')['transaction_status'].to_dict()
            missing_in_db['Current DB Status'] = missing_in_db['Receipt No.'].map(db_status_map).fillna("Missing entirely")
            
            # Apply styling and display
            st.dataframe(
                missing_in_db.style.apply(highlight_red, axis=1),
                use_container_width=True
            )
        else:
            st.success("All receipts in the Telebirr statement matched successfully in the database!")

        # Reconcile statement withdrawals against the orders marked as refunded.
        refund_query = text("""
            SELECT
                o.id AS order_id,
                o.created_at,
                o.order_amount,
                o.tip,
                o.adjusment,
                o.refunded,
                o.user_id,
                o.payment_method,
                o.payment_status AS transaction_status,
                o.refund_amount,
                o.refund_session_id
            FROM orders o
            WHERE DATE(o.refunded) BETWEEN :start_date AND :end_date
              AND o.payment_status = 'refunded'
              AND o.payment_method IN ('telebirr', 'telebirr-ussd')
        """)
        refunded_orders_df = pd.read_sql(
            refund_query,
            engine,
            params={"start_date": start_date, "end_date": end_date},
        )
        refund_report_df, statement_withdrawals = build_refund_reconciliation(
            refund_statement_df, refunded_orders_df
        )

        st.subheader("Check 3: Refund Reconciliation")
        st.write(
            "Negative statement withdrawals are compared by absolute value with the absolute "
            "database total of order_amount + tip + adjustment. Missing receipts and amount "
            "differences are highlighted in red."
        )
        if refund_report_df.empty:
            st.success("No refund withdrawals or refunded orders were found for this date range.")
        else:
            mismatch_count = int(refund_report_df["Mismatch Reason"].ne("").sum())
            if mismatch_count:
                st.warning(f"{mismatch_count} refund reconciliation issue(s) found.")
            else:
                st.success("All statement withdrawals match the refunded orders.")
            st.dataframe(
                refund_report_df.style.apply(
                    lambda row: [
                        "background-color: #ffcccc; color: #990000; font-weight: bold"
                    ] * len(row) if row["Mismatch Reason"] else [""] * len(row),
                    axis=1,
                ),
                use_container_width=True,
            )

        report_bytes = build_excel_report(
            uploaded_file.getvalue(),
            uploaded_file.name,
            refund_report_df,
            start_date,
            end_date,
        )
        output_name = f"{Path(uploaded_file.name).stem}_reconciled.xlsx"
        st.download_button(
            "Download reconciled Excel workbook",
            data=report_bytes,
            file_name=output_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except Exception as e:
        st.error(f"An error occurred during reconciliation: {e}")